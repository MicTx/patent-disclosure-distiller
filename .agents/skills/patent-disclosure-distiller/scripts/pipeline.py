"""Deterministic DisclosurePackage validation and artifact generation.

The module deliberately consumes a structured JSON package. It does not infer
technical facts or invent missing figures; the input evidence ledger remains
the single content source for Markdown, DOCX, SVG/PNG and manifest outputs.
"""

from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence, Tuple

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


EVIDENCE_TYPES = {"F", "C", "I", "A", "Q", "L"}
DELIVERY_MODES = {"FULL_DISCLOSURE", "TARGETED_SECTION", "REWRITE_AUDIT", "INTAKE_DRAFT"}
ROUTES = {"DATA_CONTROL_LOOP", "STRUCTURE_PHYSICS", "PROCESS_PROTOCOL", "HYBRID"}
FIGURE_TYPES = {"system", "flow", "sequence", "state", "structure"}
DOCX_FONT = "Hiragino Sans GB"
FIGURE_OUTPUT_MODES = {"rendered", "spec_only"}
FIGURE_GENERATION_METHODS = {"imagegen", "ppt"}
EVIDENCE_ID_RE = re.compile(r"^ev-[a-z0-9-]+$")
SECTION_ID_RE = re.compile(r"^[a-z0-9-]+$")
FIGURE_ID_RE = re.compile(r"^fig-[a-z0-9-]+$")
FULL_SECTION_IDS = {
    "delivery-scope",
    "summary",
    "technical-field",
    "background",
    "problem",
    "system-boundary",
    "solution",
    "parameters",
    "embodiments",
    "effects",
    "figure-description",
}
EVIDENCE_RE = re.compile(r"\[(F|C|I|A|Q|L)\]")


class PackageValidationError(ValueError):
    """Raised when a package cannot be safely rendered."""

    def __init__(self, errors: Sequence[str]):
        self.errors = list(errors)
        super().__init__("; ".join(self.errors))


def load_json(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise PackageValidationError(["根对象必须是 JSON object"])
    return value


def _is_nonempty_string(value: Any) -> bool:
    if not isinstance(value, str) or not value.strip():
        return False
    return not any(ord(char) < 0x20 and char not in "\t\n\r" for char in value)


def _is_allowed(value: Any, allowed: set[str]) -> bool:
    return isinstance(value, str) and value in allowed


def _string_list(value: Any) -> bool:
    return isinstance(value, list) and all(_is_nonempty_string(item) for item in value)


def _check_unique(errors: List[str], values: Iterable[Any], label: str) -> None:
    seen: Dict[Any, int] = {}
    for value in values:
        seen[value] = seen.get(value, 0) + 1
    for value, count in seen.items():
        if count > 1:
            errors.append(f"{label} 重复：{value}")


def validate_package(data: Mapping[str, Any]) -> List[str]:
    """Return deterministic, human-readable input contract errors."""

    errors: List[str] = []
    if data.get("package_version") != "1.0":
        errors.append("package_version 必须为 1.0")

    metadata = data.get("metadata")
    if not isinstance(metadata, Mapping):
        errors.append("metadata 必须是对象")
        metadata = {}
    if not _is_nonempty_string(metadata.get("title")):
        errors.append("metadata.title 不能为空")
    if not _is_allowed(metadata.get("delivery_mode"), DELIVERY_MODES):
        errors.append("metadata.delivery_mode 无效")
    if not _is_allowed(metadata.get("route"), ROUTES):
        errors.append("metadata.route 无效")
    if not _is_nonempty_string(metadata.get("release_status")):
        errors.append("metadata.release_status 不能为空")
    figure_output_mode = metadata.get("figure_output_mode", "rendered")
    if not _is_allowed(figure_output_mode, FIGURE_OUTPUT_MODES):
        errors.append("metadata.figure_output_mode 无效")

    evidence = data.get("evidence")
    if not isinstance(evidence, list):
        errors.append("evidence 必须是数组")
        evidence = []
    evidence_ids = [item.get("id") for item in evidence if isinstance(item, Mapping) and isinstance(item.get("id"), str)]
    _check_unique(errors, evidence_ids, "evidence.id")
    evidence_by_id = set(evidence_ids)
    for index, item in enumerate(evidence):
        if not isinstance(item, Mapping):
            errors.append(f"evidence[{index}] 必须是对象")
            continue
        for key in ("id", "source", "locator", "text"):
            if not _is_nonempty_string(item.get(key)):
                errors.append(f"evidence[{index}].{key} 不能为空")
        if not _is_allowed(item.get("type"), EVIDENCE_TYPES):
            errors.append(f"evidence[{index}].type 无效")
        if not EVIDENCE_ID_RE.fullmatch(str(item.get("id"))):
            errors.append(f"evidence[{index}].id 格式无效")

    sections = data.get("sections")
    if not isinstance(sections, list) or not sections:
        errors.append("sections 必须是非空数组")
        sections = []
    section_ids = [item.get("id") for item in sections if isinstance(item, Mapping) and isinstance(item.get("id"), str)]
    _check_unique(errors, section_ids, "section.id")
    section_id_set = set(section_ids)
    delivery_mode = metadata.get("delivery_mode")
    if delivery_mode == "FULL_DISCLOSURE":
        missing = sorted(FULL_SECTION_IDS - section_id_set)
        if missing:
            errors.append(f"FULL_DISCLOSURE 缺少章节：{', '.join(missing)}")
        if not isinstance(data.get("mode_contract"), Mapping):
            errors.append("FULL_DISCLOSURE 必须声明 mode_contract")
    elif delivery_mode == "TARGETED_SECTION":
        if not _string_list(data.get("requested_sections")) or not data.get("requested_sections"):
            errors.append("TARGETED_SECTION 必须声明 requested_sections")
        elif not set(section_ids).issubset(set(data.get("requested_sections"))):
            errors.append("TARGETED_SECTION 输出章节超出 requested_sections")
    elif delivery_mode == "REWRITE_AUDIT":
        audit_rows = data.get("rewrite_audit")
        if not isinstance(audit_rows, list) or not audit_rows:
            errors.append("REWRITE_AUDIT 必须声明 rewrite_audit")
        else:
            for index, row in enumerate(audit_rows):
                if not isinstance(row, Mapping) or not all(_is_nonempty_string(row.get(key)) for key in ("original", "problem", "rewrite")):
                    errors.append(f"rewrite_audit[{index}] 必须有 original/problem/rewrite")
    elif delivery_mode == "INTAKE_DRAFT":
        if not data.get("questions"):
            errors.append("INTAKE_DRAFT 必须至少有一个待确认问题")

    table_map: Dict[str, Mapping[str, Any]] = {}
    tables = data.get("tables", [])
    if not isinstance(tables, list):
        errors.append("tables 必须是数组")
        tables = []
    for index, table in enumerate(tables):
        if not isinstance(table, Mapping):
            errors.append(f"tables[{index}] 必须是对象")
            continue
        table_id = table.get("id")
        if not _is_nonempty_string(table_id):
            errors.append(f"tables[{index}].id 不能为空")
        elif table_id in table_map:
            errors.append(f"table.id 重复：{table_id}")
        else:
            table_map[str(table_id)] = table
        if not _string_list(table.get("columns")) or not table.get("columns"):
            errors.append(f"tables[{index}].columns 必须为非空数组")
        if not isinstance(table.get("rows"), list):
            errors.append(f"tables[{index}].rows 必须为数组")
        else:
            for row_index, row in enumerate(table.get("rows", [])):
                if not isinstance(row, (list, Mapping)):
                    errors.append(f"tables[{index}].rows[{row_index}] 必须是数组或对象")
                elif isinstance(row, list) and len(row) != len(table.get("columns", [])):
                    errors.append(f"tables[{index}].rows[{row_index}] 列数与 columns 不一致")
                elif isinstance(row, list) and any(not _is_nonempty_string(str(value)) for value in row):
                    errors.append(f"tables[{index}].rows[{row_index}] 含空值或非法文本")
                elif isinstance(row, Mapping) and any(not _is_nonempty_string(str(value)) for value in row.values()):
                    errors.append(f"tables[{index}].rows[{row_index}] 含空值或非法文本")
        if "evidence_note" in table and not _is_nonempty_string(table.get("evidence_note")):
            errors.append(f"tables[{index}].evidence_note 文本非法")
        table_evidence = table.get("evidence", [])
        if delivery_mode == "FULL_DISCLOSURE" and (not _string_list(table_evidence) or not table_evidence):
            errors.append(f"FULL_DISCLOSURE tables[{index}] 必须声明 evidence")
        for evidence_id in table_evidence if isinstance(table_evidence, list) else []:
            if not isinstance(evidence_id, str) or evidence_id not in evidence_by_id:
                errors.append(f"tables[{index}] 引用不存在证据：{evidence_id}")

    figures = data.get("figures")
    if not isinstance(figures, list):
        errors.append("figures 必须是数组")
        figures = []
    figure_ids = [item.get("id") for item in figures if isinstance(item, Mapping) and isinstance(item.get("id"), str)]
    _check_unique(errors, figure_ids, "figure.id")
    figure_id_set = set(figure_ids)
    if delivery_mode == "FULL_DISCLOSURE" and len(figures) < 3:
        errors.append("FULL_DISCLOSURE 至少需要 3 张附图")
    if delivery_mode == "FULL_DISCLOSURE":
        figure_types = {figure.get("type") for figure in figures if isinstance(figure, Mapping) and isinstance(figure.get("type"), str)}
        for required_type in ("system", "flow", "sequence"):
            if required_type not in figure_types:
                errors.append(f"FULL_DISCLOSURE 缺少附图类型：{required_type}")

    for index, figure in enumerate(figures):
        if not isinstance(figure, Mapping):
            errors.append(f"figures[{index}] 必须是对象")
            continue
        for key in ("id", "title", "caption"):
            if not _is_nonempty_string(figure.get(key)):
                errors.append(f"figures[{index}].{key} 不能为空")
        if not FIGURE_ID_RE.fullmatch(str(figure.get("id"))):
            errors.append(f"figures[{index}].id 格式无效")
        if not _is_allowed(figure.get("type"), FIGURE_TYPES):
            errors.append(f"figures[{index}].type 无效")
        if delivery_mode == "FULL_DISCLOSURE" and not _is_nonempty_string(figure.get("source_image")):
            errors.append(f"FULL_DISCLOSURE figures[{index}] 必须声明由 imagegen 或 PPT 导出的 source_image")
        if delivery_mode == "FULL_DISCLOSURE" and not _is_allowed(figure.get("generation_method"), FIGURE_GENERATION_METHODS):
            errors.append(f"FULL_DISCLOSURE figures[{index}] 必须声明 generation_method=imagegen 或 ppt")
        if delivery_mode == "FULL_DISCLOSURE" and not _is_nonempty_string(figure.get("prompt")):
            errors.append(f"FULL_DISCLOSURE figures[{index}] 必须声明 imagegen/PPT 提示词或绘制说明")
        if not _string_list(figure.get("labels")) or not figure.get("labels"):
            errors.append(f"figures[{index}].labels 必须为非空数组")
        if not _string_list(figure.get("used_by")) or not figure.get("used_by"):
            errors.append(f"figures[{index}].used_by 必须为非空数组")
        for used_by in figure.get("used_by", []) if isinstance(figure.get("used_by"), list) else []:
            if not isinstance(used_by, str) or used_by not in section_id_set:
                errors.append(f"figures[{index}].used_by 引用不存在章节：{used_by}")
        figure_evidence = figure.get("evidence", [])
        if not _string_list(figure_evidence) or not figure_evidence:
            errors.append(f"figures[{index}].evidence 不能为空")
        for evidence_id in figure_evidence if isinstance(figure_evidence, list) else []:
            if not isinstance(evidence_id, str) or evidence_id not in evidence_by_id:
                errors.append(f"figures[{index}] 引用不存在证据：{evidence_id}")

    for section_index, section in enumerate(sections):
        if not isinstance(section, Mapping):
            errors.append(f"sections[{section_index}] 必须是对象")
            continue
        if not _is_nonempty_string(section.get("id")) or not SECTION_ID_RE.fullmatch(str(section.get("id"))) or not _is_nonempty_string(section.get("title")):
            errors.append(f"sections[{section_index}] 必须有 id/title")
        level = section.get("level")
        if not isinstance(level, int) or level < 1 or level > 3:
            errors.append(f"sections[{section_index}].level 必须为 1-3")
        blocks = section.get("blocks")
        if not isinstance(blocks, list) or not blocks:
            errors.append(f"sections[{section_index}].blocks 必须为非空数组")
            continue
        for block_index, block in enumerate(blocks):
            if not isinstance(block, Mapping):
                errors.append(f"sections[{section_index}].blocks[{block_index}] 必须是对象")
                continue
            block_type = block.get("type", "paragraph")
            if block_type in {"paragraph", "bullet"}:
                if not _is_nonempty_string(block.get("text")):
                    errors.append(f"sections[{section_index}].blocks[{block_index}].text 不能为空")
                refs = block.get("evidence", [])
                _validate_block_evidence(errors, data, block, f"sections[{section_index}].blocks[{block_index}]")
                for evidence_id in refs if isinstance(refs, list) else []:
                    if not isinstance(evidence_id, str) or evidence_id not in evidence_by_id:
                        errors.append(f"正文块引用不存在证据：{evidence_id}")
            elif block_type == "table_ref":
                if block.get("table_id") not in table_map:
                    errors.append(f"正文块引用不存在表格：{block.get('table_id')}")
            elif block_type == "figure_ref":
                if block.get("figure_id") not in figure_id_set:
                    errors.append(f"正文块引用不存在附图：{block.get('figure_id')}")
            else:
                errors.append(f"不支持的 block.type：{block_type}")

    # Validate graph primitives before any renderer touches them.
    for figure_index, figure in enumerate(figures):
        if not isinstance(figure, Mapping):
            continue
        nodes = figure.get("nodes", [])
        if not isinstance(nodes, list) or any(not isinstance(node, Mapping) for node in nodes):
            errors.append(f"figures[{figure_index}].nodes 必须是对象数组")
            nodes = []
        node_ids = [node.get("id") for node in nodes]
        if not node_ids or any(not _is_nonempty_string(node_id) for node_id in node_ids) or len(set(node_ids)) != len(node_ids):
            errors.append(f"figures[{figure_index}].nodes.id 必须唯一且非空")
        for node_index, node in enumerate(nodes):
            if not _is_nonempty_string(node.get("label")) or ("role" in node and not _is_nonempty_string(node.get("role"))):
                errors.append(f"figures[{figure_index}].nodes[{node_index}] 文本非法")
        edges = figure.get("edges", [])
        if not isinstance(edges, list) or any(not isinstance(edge, Mapping) for edge in edges):
            errors.append(f"figures[{figure_index}].edges 必须是对象数组")
            edges = []
        if not edges:
            errors.append(f"figures[{figure_index}].edges 不能为空")
        for edge_index, edge in enumerate(edges):
            source = edge.get("from", edge.get("source"))
            target = edge.get("to", edge.get("target"))
            if source not in node_ids or target not in node_ids:
                errors.append(f"figures[{figure_index}].edges[{edge_index}] 引用不存在节点")
            if "label" in edge and not _is_nonempty_string(edge.get("label")):
                errors.append(f"figures[{figure_index}].edges[{edge_index}].label 文本非法")

    referenced_figures = {
        block.get("figure_id")
        for section in sections
        if isinstance(section, Mapping)
        for block in section.get("blocks", [])
        if isinstance(block, Mapping) and block.get("type") == "figure_ref"
    }
    for figure_id in figure_id_set - referenced_figures:
        errors.append(f"附图未被正文 figure_ref 使用：{figure_id}")

    questions = data.get("questions")
    if not isinstance(questions, list):
        errors.append("questions 必须是数组")
    else:
        for index, question in enumerate(questions):
            if not isinstance(question, Mapping) or not _is_nonempty_string(question.get("id")) or not _is_nonempty_string(question.get("text")):
                errors.append(f"questions[{index}] 必须有 id/text")

    audit = data.get("audit")
    if not isinstance(audit, Mapping):
        errors.append("audit 必须是对象")
    else:
        for key in ("p0", "p1", "p2"):
            if not isinstance(audit.get(key), list):
                errors.append(f"audit.{key} 必须是数组")
        if audit.get("p0"):
            errors.append("audit.p0 非空，存在未关闭的阻断问题")
    return errors


def ensure_valid(data: Mapping[str, Any]) -> None:
    errors = validate_package(data)
    if errors:
        raise PackageValidationError(errors)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _evidence_suffix(data: Mapping[str, Any], refs: Sequence[str], primary: str | None = None) -> str:
    if primary:
        return f" [{primary}]"
    by_id = {item.get("id"): item for item in data.get("evidence", []) if isinstance(item, Mapping) and isinstance(item.get("id"), str)}
    tags = []
    for ref in refs:
        tag = by_id.get(ref, {}).get("type")
        if tag:
            tags.append(f"[{tag}]")
    return " " + " ".join(tags) if tags else ""


def _validate_block_evidence(errors: List[str], data: Mapping[str, Any], block: Mapping[str, Any], location: str) -> None:
    refs = block.get("evidence", [])
    if not _string_list(refs) or not refs:
        errors.append(f"{location} 必须声明 evidence")
        return
    by_id = {item.get("id"): item for item in data.get("evidence", []) if isinstance(item, Mapping) and isinstance(item.get("id"), str)}
    types = {by_id[ref].get("type") for ref in refs if ref in by_id and isinstance(by_id[ref].get("type"), str)}
    primary = block.get("evidence_primary")
    if len(types) > 1 and (not isinstance(primary, str) or primary not in types):
        errors.append(f"{location} 包含多个证据类型，必须声明 evidence_primary")
    if primary is not None and (not isinstance(primary, str) or primary not in types):
        errors.append(f"{location}.evidence_primary 不在 evidence 引用中")
    inline_tags = set(EVIDENCE_RE.findall(str(block.get("text", ""))))
    if inline_tags and (len(inline_tags) > 1 or (isinstance(primary, str) and primary not in inline_tags)):
        errors.append(f"{location} 正文内嵌证据标签与 evidence_primary 冲突")


def _figure_number_map(data: Mapping[str, Any]) -> Dict[str, int]:
    return {figure["id"]: index for index, figure in enumerate(data.get("figures", []), 1)}


def render_markdown(data: Mapping[str, Any], output: Path) -> None:
    figure_numbers = _figure_number_map(data)
    lines = [f"# {data['metadata']['title']}", "", f"> 交付状态：{data['metadata']['release_status']}", "", "## 证据状态", ""]
    lines.extend([
        "`[F]` 材料事实；`[C]` 实际核验；`[I]` 推导；`[A]` 工作假设；`[Q]` 待确认；`[L]` 法律审查。",
        "",
    ])
    for section in data["sections"]:
        lines.append("#" * (section["level"] + 1) + " " + section["title"])
        lines.append("")
        for block in section["blocks"]:
            block_type = block.get("type", "paragraph")
            if block_type == "paragraph":
                lines.append(block["text"] + _evidence_suffix(data, block.get("evidence", []), block.get("evidence_primary")))
                lines.append("")
            elif block_type == "bullet":
                lines.append("- " + block["text"] + _evidence_suffix(data, block.get("evidence", []), block.get("evidence_primary")))
            elif block_type == "table_ref":
                table = next(item for item in data.get("tables", []) if item["id"] == block["table_id"])
                lines.extend(_markdown_table(table, data))
                lines.append("")
            elif block_type == "figure_ref":
                figure = next(item for item in data["figures"] if item["id"] == block["figure_id"])
                number = figure_numbers[figure["id"]]
                if data["metadata"].get("figure_output_mode", "rendered") == "spec_only":
                    lines.extend([
                        f"图 {number}：{figure['title']}（附图方案，实际图由 {figure['generation_method']} 生成）",
                        f"- 生成方式：`{figure['generation_method']}`",
                        f"- 生成说明：{figure['prompt']}",
                        f"- 底图来源：`{figure['source_image']}`",
                        f"- 图注：{figure['caption']} {_evidence_suffix(data, figure.get('evidence', []))}",
                        "",
                    ])
                else:
                    lines.extend([
                        f"图 {number}：{figure['title']}",
                        f"![图 {number}：{figure['title']}](figures/{figure['id']}.png)",
                        f"{figure['caption']} {_evidence_suffix(data, figure.get('evidence', []))}",
                        "",
                    ])
        lines.append("")
    if data["metadata"].get("delivery_mode") == "REWRITE_AUDIT":
        lines.extend(["## 改写审计", "", "| 原句 | 问题 | 改写 |", "| --- | --- | --- |"])
        for row in data.get("rewrite_audit", []):
            lines.append(f"| {row['original']} | {row['problem']} | {row['rewrite']} |")
        lines.append("")
    lines.extend(["## 待确认信息", ""])
    for question in data.get("questions", []):
        lines.append(f"- **{question['id']}** {question['text']}（{question.get('status', 'open')}）")
    lines.extend(["", "## 自审结果", "", f"- P0：{len(data['audit']['p0'])}", f"- P1：{len(data['audit']['p1'])}", f"- P2：{len(data['audit']['p2'])}", ""])
    output.write_text("\n".join(lines), encoding="utf-8")


def _markdown_table(table: Mapping[str, Any], data: Mapping[str, Any] | None = None) -> List[str]:
    columns = [str(value) for value in table["columns"]]
    rows = table["rows"]
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join("---" for _ in columns) + " |"]
    for row in rows:
        values = row if isinstance(row, list) else [row.get(column, "") for column in columns]
        lines.append("| " + " | ".join(str(value) for value in values) + " |")
    if data is not None:
        lines.append("")
        lines.append("表格证据：" + _evidence_suffix(data, table.get("evidence", [])))
    return lines


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    paragraph.add_run(" 页")


def _configure_docx(document: Document, title: str) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = DOCX_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 16), ("Heading 3", 14)):
        style = styles[name]
        style.font.name = DOCX_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = __import__("docx").shared.RGBColor.from_string("000000")
        style.paragraph_format.keep_with_next = True
    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.base_style = normal
    caption = styles["Figure Caption"]
    caption.font.name = DOCX_FONT
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    _add_page_number(footer)


def _find_reference_docx(input_path: Path, reference: str | None) -> Path | None:
    if not reference:
        return None
    requested = Path(reference)
    candidates = [input_path.parent / requested, Path.cwd() / requested]
    candidates.extend(parent / requested for parent in input_path.parents)
    for candidate in candidates:
        if candidate.is_file() and candidate.suffix.lower() == ".docx":
            return candidate.resolve()
    return None


def _clear_document_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _outline_level(paragraph: Any) -> int | None:
    ppr = paragraph._p.pPr
    if ppr is None or ppr.outlineLvl is None:
        return None
    try:
        return int(ppr.outlineLvl.get(qn("w:val")))
    except (TypeError, ValueError):
        return None


def _sample_donors(document: Document) -> Dict[str, Any]:
    paragraphs = list(document.paragraphs)
    body = next((paragraph for paragraph in paragraphs if _outline_level(paragraph) is None and paragraph.runs and not any(run.bold for run in paragraph.runs)), paragraphs[0])
    by_level = {level: next((paragraph for paragraph in paragraphs if _outline_level(paragraph) == level), body) for level in (0, 1, 2)}
    bullet = next((paragraph for paragraph in paragraphs if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None), body)
    return {"title": paragraphs[0], "body": body, "bullet": bullet, "heading": by_level}


def _add_sample_paragraph(document: Document, donor: Any, text: str) -> Any:
    paragraph = document.add_paragraph()
    for name, value in donor._p.attrib.items():
        paragraph._p.set(name, value)
    if donor._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(donor._p.pPr))
    run = paragraph.add_run(text)
    if donor.runs and donor.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(donor.runs[0]._r.rPr))
    return paragraph


def render_docx_sample_fidelity(data: Mapping[str, Any], output: Path, figure_dir: Path, reference_docx: Path, exact: bool = False) -> None:
    if exact:
        shutil.copy2(reference_docx, output)
        return
    document = Document(str(reference_docx))
    donors = _sample_donors(document)
    _clear_document_body(document)
    _add_sample_paragraph(document, donors["title"], data["metadata"]["title"])
    _add_sample_paragraph(document, donors["body"], f"技术交底书底稿｜{data['metadata']['release_status']}")
    for section in data["sections"]:
        level = min(max(int(section["level"]) - 1, 0), 2)
        _add_sample_paragraph(document, donors["heading"][level], section["title"])
        for block in section["blocks"]:
            block_type = block.get("type", "paragraph")
            if block_type in {"paragraph", "bullet"}:
                text = block["text"] + _evidence_suffix(data, block.get("evidence", []), block.get("evidence_primary"))
                _add_sample_paragraph(document, donors["bullet"] if block_type == "bullet" else donors["body"], text)
            elif block_type == "table_ref":
                table = next(item for item in data.get("tables", []) if item["id"] == block["table_id"])
                _add_docx_table(document, table, data)
            elif block_type == "figure_ref":
                figure = next(item for item in data["figures"] if item["id"] == block["figure_id"])
                if data["metadata"].get("figure_output_mode", "rendered") == "spec_only":
                    _add_sample_paragraph(document, donors["body"], f"附图方案：{figure['title']}；生成方式：{figure['generation_method']}；说明：{figure['prompt']}；底图来源：{figure['source_image']}。")
                else:
                    document.add_picture(str(figure_dir / f"{figure['id']}.png"), width=Cm(15.5))
                    _add_sample_paragraph(document, donors["body"], figure["caption"] + _evidence_suffix(data, figure.get("evidence", [])))
    _add_sample_paragraph(document, donors["heading"][0], "待确认信息")
    for question in data.get("questions", []):
        _add_sample_paragraph(document, donors["bullet"], f"{question['id']}：{question['text']}（{question.get('status', 'open')}）")
    _add_sample_paragraph(document, donors["heading"][0], "自审结果")
    for key in ("p0", "p1", "p2"):
        _add_sample_paragraph(document, donors["bullet"], f"{key.upper()}：{len(data['audit'][key])}")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    _normalize_zip_timestamps(output)


def render_docx(data: Mapping[str, Any], output: Path, figure_dir: Path, reference_docx: Path | None = None) -> None:
    profile = data["metadata"].get("docx_profile", "native")
    if profile in {"sample_fidelity", "sample_exact"}:
        if reference_docx is None:
            raise PackageValidationError(["sample_fidelity/sample_exact 必须提供 reference_docx"])
        render_docx_sample_fidelity(data, output, figure_dir, reference_docx, exact=profile == "sample_exact")
        return
    document = Document()
    _configure_docx(document, data["metadata"]["title"])
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(data["metadata"]["title"])
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"技术交底书底稿｜{data['metadata']['release_status']}").bold = True
    document.add_paragraph("本文件由结构化 DisclosurePackage 生成；未关闭的证据、实验和法律事项保留在正文与自审结果中。")
    figure_numbers = _figure_number_map(data)
    for section in data["sections"]:
        paragraph = document.add_paragraph(style=f"Heading {section['level']}")
        paragraph.add_run(section["title"])
        for block in section["blocks"]:
            block_type = block.get("type", "paragraph")
            if block_type == "paragraph":
                document.add_paragraph(block["text"] + _evidence_suffix(data, block.get("evidence", []), block.get("evidence_primary")))
            elif block_type == "bullet":
                document.add_paragraph(block["text"] + _evidence_suffix(data, block.get("evidence", []), block.get("evidence_primary")), style="List Bullet")
            elif block_type == "table_ref":
                table = next(item for item in data.get("tables", []) if item["id"] == block["table_id"])
                _add_docx_table(document, table, data)
            elif block_type == "figure_ref":
                figure = next(item for item in data["figures"] if item["id"] == block["figure_id"])
                if data["metadata"].get("figure_output_mode", "rendered") == "spec_only":
                    document.add_paragraph(f"附图方案：{figure['title']}；生成方式：{figure['generation_method']}；说明：{figure['prompt']}；底图来源：{figure['source_image']}。")
                else:
                    number = figure_numbers[figure["id"]]
                    document.add_picture(str(figure_dir / f"{figure['id']}.png"), width=Cm(15.5))
                    caption = document.add_paragraph(style="Figure Caption")
                    caption.add_run(f"图 {number}  {figure['title']}")
                    document.add_paragraph(figure["caption"] + _evidence_suffix(data, figure.get("evidence", [])))
    document.add_paragraph("待确认信息", style="Heading 1")
    if data["metadata"].get("delivery_mode") == "REWRITE_AUDIT":
        document.add_paragraph("改写审计", style="Heading 1")
        audit_table = document.add_table(rows=1, cols=3)
        try:
            audit_table.style = "Table Grid"
        except KeyError:
            pass
        for index, header in enumerate(("原句", "问题", "改写")):
            audit_table.rows[0].cells[index].text = header
            _set_cell_shading(audit_table.rows[0].cells[index], "DCEAF2")
        for row in data.get("rewrite_audit", []):
            cells = audit_table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = row["original"], row["problem"], row["rewrite"]
    for question in data.get("questions", []):
        document.add_paragraph(f"{question['id']}：{question['text']}（{question.get('status', 'open')}）", style="List Bullet")
    document.add_paragraph("自审结果", style="Heading 1")
    for key in ("p0", "p1", "p2"):
        document.add_paragraph(f"{key.upper()}：{len(data['audit'][key])}", style="List Bullet")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    _normalize_zip_timestamps(output)


def _normalize_zip_timestamps(path: Path) -> None:
    """Make python-docx ZIP output reproducible across invocations."""
    with zipfile.ZipFile(path, "r") as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with tempfile.NamedTemporaryFile(dir=str(path.parent), suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for info, payload in entries:
                normalized = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                normalized.compress_type = zipfile.ZIP_DEFLATED
                normalized.external_attr = info.external_attr
                normalized.create_system = info.create_system
                target.writestr(normalized, payload)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _add_docx_table(document: Document, table_spec: Mapping[str, Any], data: Mapping[str, Any]) -> None:
    columns = [str(value) for value in table_spec["columns"]]
    table = document.add_table(rows=1, cols=len(columns))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for index, column in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = column
        _set_cell_shading(cell, "DCEAF2")
    for row in table_spec["rows"]:
        values = row if isinstance(row, list) else [row.get(column, "") for column in columns]
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    evidence = _evidence_suffix(data, table_spec.get("evidence", []))
    document.add_paragraph(table_spec.get("evidence_note", "表格数据来自输入包。") + evidence)


def _font(size: int) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size)
            except OSError:
                continue
    return ImageFont.load_default()


def _graph_layout(figure: Mapping[str, Any]) -> Tuple[List[Mapping[str, Any]], List[Mapping[str, Any]]]:
    nodes = figure.get("nodes") or [{"id": f"n{index}", "label": label} for index, label in enumerate(figure.get("labels", []), 1)]
    edges = figure.get("edges") or []
    return list(nodes), list(edges)


def _svg_for_figure(figure: Mapping[str, Any], number: int) -> str:
    if figure.get("type") == "sequence":
        return _sequence_svg_for_figure(figure, number)
    nodes, edges = _graph_layout(figure)
    width, height = 1200, max(700, 180 + 150 * ((len(nodes) + 3) // 4))
    positions: Dict[str, Tuple[int, int]] = {}
    for index, node in enumerate(nodes):
        node_id = str(node.get("id", f"n{index}"))
        col, row = index % 4, index // 4
        positions[node_id] = (80 + col * 285, 150 + row * 180)
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#245B73"/></marker></defs>',
        f'<rect width="100%" height="100%" fill="#F8FAFC"/><text x="50" y="55" font-size="30" font-family="Arial, sans-serif" font-weight="700" fill="#16324F">图 {number}  {html.escape(str(figure["title"]))}</text>',
    ]
    for edge in edges:
        source = str(edge.get("from", edge.get("source", "")))
        target = str(edge.get("to", edge.get("target", "")))
        if source not in positions or target not in positions:
            continue
        x1, y1 = positions[source]
        x2, y2 = positions[target]
        parts.append(f'<path d="M{x1+210},{y1+45} L{x2},{y2+45}" stroke="#245B73" stroke-width="3" fill="none" marker-end="url(#arrow)"/>')
        label = edge.get("label")
        if label:
            parts.append(f'<text x="{(x1+x2)//2}" y="{(y1+y2)//2+35}" font-size="18" font-family="Arial, sans-serif" fill="#334155">{html.escape(str(label))}</text>')
    for index, node in enumerate(nodes):
        node_id = str(node.get("id", f"n{index}"))
        x, y = positions[node_id]
        label = html.escape(str(node.get("label", node_id)))
        role = html.escape(str(node.get("role", "")))
        parts.append(f'<rect x="{x}" y="{y}" width="210" height="90" rx="12" fill="#E7F1F5" stroke="#245B73" stroke-width="3"/>')
        parts.append(f'<text x="{x+105}" y="{y+38}" text-anchor="middle" font-size="20" font-family="Arial, sans-serif" font-weight="700" fill="#16324F">{label}</text>')
        if role:
            parts.append(f'<text x="{x+105}" y="{y+67}" text-anchor="middle" font-size="15" font-family="Arial, sans-serif" fill="#475569">{role}</text>')
    labels = "；".join(str(value) for value in figure.get("labels", []))
    parts.append(f'<text x="50" y="{height-40}" font-size="16" font-family="Arial, sans-serif" fill="#475569">关键标识：{html.escape(labels)}</text>')
    parts.append("</svg>")
    return "".join(parts)


def _sequence_svg_for_figure(figure: Mapping[str, Any], number: int) -> str:
    nodes, edges = _graph_layout(figure)
    width, height = 1200, max(700, 330 + 75 * len(edges))
    positions = {str(node.get("id")): 90 + index * (width - 180) // max(1, len(nodes) - 1) for index, node in enumerate(nodes)}
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#245B73"/></marker></defs>',
        f'<rect width="100%" height="100%" fill="#F8FAFC"/><text x="50" y="55" font-size="30" font-family="Arial, sans-serif" font-weight="700" fill="#16324F">图 {number}  {html.escape(str(figure["title"]))}</text>',
    ]
    for node in nodes:
        node_id = str(node.get("id"))
        x = positions[node_id]
        label = html.escape(str(node.get("label", node_id)))
        role = html.escape(str(node.get("role", "")))
        parts.append(f'<rect x="{x-85}" y="100" width="170" height="70" rx="10" fill="#E7F1F5" stroke="#245B73" stroke-width="3"/>')
        parts.append(f'<text x="{x}" y="132" text-anchor="middle" font-size="19" font-family="Arial, sans-serif" font-weight="700" fill="#16324F">{label}</text>')
        parts.append(f'<text x="{x}" y="155" text-anchor="middle" font-size="13" font-family="Arial, sans-serif" fill="#475569">{role}</text>')
        parts.append(f'<path d="M{x},{175} L{x},{height-80}" stroke="#94A3B8" stroke-width="2" stroke-dasharray="8 8"/>')
    for index, edge in enumerate(edges):
        source = str(edge.get("from", edge.get("source", "")))
        target = str(edge.get("to", edge.get("target", "")))
        if source not in positions or target not in positions:
            continue
        y = 235 + index * 65
        x1, x2 = positions[source], positions[target]
        parts.append(f'<path d="M{x1},{y} L{x2},{y}" stroke="#245B73" stroke-width="3" fill="none" marker-end="url(#arrow)"/>')
        if edge.get("label"):
            parts.append(f'<text x="{(x1+x2)//2}" y="{y-10}" text-anchor="middle" font-size="16" font-family="Arial, sans-serif" fill="#334155">{html.escape(str(edge["label"]))}</text>')
    labels = "；".join(str(value) for value in figure.get("labels", []))
    parts.append(f'<text x="50" y="{height-35}" font-size="16" font-family="Arial, sans-serif" fill="#475569">关键标识：{html.escape(labels)}</text>')
    parts.append("</svg>")
    return "".join(parts)


def _png_for_figure(figure: Mapping[str, Any], number: int, output: Path) -> None:
    if figure.get("type") == "sequence":
        _sequence_png_for_figure(figure, number, output)
        return
    nodes, edges = _graph_layout(figure)
    width, height = 1200, max(700, 180 + 150 * ((len(nodes) + 3) // 4))
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font, node_font, small_font = _font(30), _font(20), _font(15)
    draw.text((50, 25), f"图 {number}  {figure['title']}", fill="#16324F", font=title_font)
    positions: Dict[str, Tuple[int, int]] = {}
    for index, node in enumerate(nodes):
        node_id = str(node.get("id", f"n{index}"))
        col, row = index % 4, index // 4
        positions[node_id] = (80 + col * 285, 150 + row * 180)
    for edge in edges:
        source, target = str(edge.get("from", edge.get("source", ""))), str(edge.get("to", edge.get("target", "")))
        if source not in positions or target not in positions:
            continue
        x1, y1 = positions[source]
        x2, y2 = positions[target]
        draw.line((x1 + 210, y1 + 45, x2, y2 + 45), fill="#245B73", width=4)
        draw.polygon(_arrow_polygon(x2, y2 + 45, x1, y1 + 45, 14), fill="#245B73")
        if edge.get("label"):
            draw.text(((x1 + x2) // 2, (y1 + y2) // 2 + 30), str(edge["label"]), fill="#334155", font=small_font)
    for index, node in enumerate(nodes):
        node_id = str(node.get("id", f"n{index}"))
        x, y = positions[node_id]
        draw.rounded_rectangle((x, y, x + 210, y + 90), radius=12, fill="#E7F1F5", outline="#245B73", width=3)
        label = str(node.get("label", node_id))
        role = str(node.get("role", ""))
        draw.text((x + 105, y + 20), label, anchor="ma", fill="#16324F", font=node_font)
        if role:
            draw.text((x + 105, y + 55), role, anchor="ma", fill="#475569", font=small_font)
    draw.text((50, height - 45), "关键标识：" + "；".join(str(value) for value in figure.get("labels", [])), fill="#475569", font=small_font)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, format="PNG", optimize=True)


def _sequence_png_for_figure(figure: Mapping[str, Any], number: int, output: Path) -> None:
    nodes, edges = _graph_layout(figure)
    width, height = 1200, max(700, 330 + 75 * len(edges))
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    title_font, node_font, small_font = _font(30), _font(19), _font(14)
    draw.text((50, 25), f"图 {number}  {figure['title']}", fill="#16324F", font=title_font)
    positions = {str(node.get("id")): 90 + index * (width - 180) // max(1, len(nodes) - 1) for index, node in enumerate(nodes)}
    for node in nodes:
        node_id = str(node.get("id"))
        x = positions[node_id]
        draw.rounded_rectangle((x - 85, 100, x + 85, 170), radius=10, fill="#E7F1F5", outline="#245B73", width=3)
        draw.text((x, 115), str(node.get("label", node_id)), anchor="ma", fill="#16324F", font=node_font)
        draw.text((x, 143), str(node.get("role", "")), anchor="ma", fill="#475569", font=small_font)
        for y in range(180, height - 80, 16):
            draw.line((x, y, x, min(y + 8, height - 80)), fill="#94A3B8", width=2)
    for index, edge in enumerate(edges):
        source = str(edge.get("from", edge.get("source", "")))
        target = str(edge.get("to", edge.get("target", "")))
        if source not in positions or target not in positions:
            continue
        y = 235 + index * 65
        x1, x2 = positions[source], positions[target]
        draw.line((x1, y, x2, y), fill="#245B73", width=4)
        draw.polygon(_arrow_polygon(x2, y, x1, y, 14), fill="#245B73")
        if edge.get("label"):
            draw.text(((x1 + x2) // 2, y - 25), str(edge["label"]), anchor="mm", fill="#334155", font=small_font)
    draw.text((50, height - 45), "关键标识：" + "；".join(str(value) for value in figure.get("labels", [])), fill="#475569", font=small_font)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, format="PNG", optimize=True)


def _arrow_polygon(tip_x: int, tip_y: int, source_x: int, source_y: int, size: int) -> List[Tuple[int, int]]:
    """Return a triangle whose tip follows the SVG marker direction."""
    dx, dy = tip_x - source_x, tip_y - source_y
    length = max((dx * dx + dy * dy) ** 0.5, 1.0)
    ux, uy = dx / length, dy / length
    base_x, base_y = tip_x - ux * size, tip_y - uy * size
    px, py = -uy * size / 2, ux * size / 2
    return [(tip_x, tip_y), (round(base_x + px), round(base_y + py)), (round(base_x - px), round(base_y - py))]


def render_figures(data: Mapping[str, Any], figure_dir: Path, source_root: Path | None = None) -> List[Dict[str, Any]]:
    figure_dir.mkdir(parents=True, exist_ok=True)
    if figure_dir.is_symlink():
        raise PackageValidationError([f"拒绝向 symlink 图目录写入：{figure_dir}"])
    rendered = []
    for number, figure in enumerate(data.get("figures", []), 1):
        svg_path = figure_dir / f"{figure['id']}.svg"
        png_path = figure_dir / f"{figure['id']}.png"
        for path in (svg_path, png_path):
            if path.exists() and path.is_symlink():
                raise PackageValidationError([f"拒绝覆盖 symlink 附图：{path}"])
        svg_path.write_text(_svg_for_figure(figure, number), encoding="utf-8")
        source_png = None
        source_image = figure.get("source_image")
        if source_image and source_root is not None:
            source_path = (source_root / source_image).resolve()
            source_root_resolved = source_root.resolve()
            if source_root_resolved not in source_path.parents or not source_path.is_file():
                raise PackageValidationError([f"生成图片底图不存在或越界：{source_image}"])
            source_png_path = figure_dir / f"{figure['id']}-source.png"
            shutil.copy2(source_path, source_png_path)
            _imagegen_composite(figure, source_path, number, png_path)
            source_png = str(source_png_path.relative_to(figure_dir.parent))
        else:
            _png_for_figure(figure, number, png_path)
        rendered.append({
            "id": figure["id"],
            "number": number,
            "type": figure["type"],
            "title": figure["title"],
            "caption": figure["caption"],
            "svg": str(svg_path.relative_to(figure_dir.parent)),
            "png": str(png_path.relative_to(figure_dir.parent)),
            "source_image": source_image,
            "source_png": source_png,
            "render_mode": "imagegen-overlay" if source_image else "code-native",
            "used_by": list(figure["used_by"]),
            "evidence": list(figure["evidence"]),
        })
    return rendered


def render_figure_specs(data: Mapping[str, Any], output_dir: Path) -> List[Dict[str, Any]]:
    """Write only auditable figure plans; actual bitmaps stay out of the template."""
    plan_path = output_dir / "figure-plan.md"
    lines = ["# 附图生成方案", "", "本交付物只保留 imagegen/PPT 生成规则、底图来源和确定性标注契约，不嵌入实际图片。", ""]
    entries = []
    for number, figure in enumerate(data.get("figures", []), 1):
        edge_text = "；".join(f"{edge['from']}→{edge['to']}({edge.get('label', '')})" for edge in figure['edges'])
        lines.extend([
            f"## 图 {number}：{figure['title']}",
            f"- figure_id：`{figure['id']}`",
            f"- 生成方式：`{figure['generation_method']}`",
            f"- 底图来源：`{figure['source_image']}`",
            f"- 生成提示/绘制说明：{figure['prompt']}",
            f"- 图注：{figure['caption']}",
            f"- 节点：{'；'.join(str(node['label']) for node in figure['nodes'])}",
            f"- 边：{edge_text}",
            f"- 正文使用：{'、'.join(figure['used_by'])}",
            "",
        ])
        entries.append({
            "id": figure["id"],
            "number": number,
            "type": figure["type"],
            "title": figure["title"],
            "caption": figure["caption"],
            "generation_method": figure["generation_method"],
            "prompt": figure["prompt"],
            "source_image": figure["source_image"],
            "render_mode": "spec_only",
            "used_by": list(figure["used_by"]),
            "evidence": list(figure["evidence"]),
        })
    plan_path.write_text("\n".join(lines), encoding="utf-8")
    return entries


def _imagegen_composite(figure: Mapping[str, Any], source_path: Path, number: int, output: Path) -> None:
    """Overlay deterministic labels onto a generated, text-free diagram base."""
    image = Image.open(source_path).convert("RGB")
    draw = ImageDraw.Draw(image)
    width, height = image.size
    title_font, node_font, small_font = _font(max(24, width // 65)), _font(max(18, width // 78)), _font(max(13, width // 105))
    draw.text((35, 22), f"图 {number}  {figure['title']}", fill="#16324F", font=title_font)
    nodes = list(figure.get("nodes", []))
    node_count = max(1, len(nodes))
    kind = figure.get("type")
    if kind == "sequence":
        y = max(100, int(height * 0.10))
        for index, node in enumerate(nodes):
            x = int((index + 0.5) * width / node_count)
            draw.text((x, y), str(node.get("label", node.get("id", ""))), anchor="ma", fill="#16324F", font=node_font)
            if node.get("role"):
                draw.text((x, y + node_font.size + 4), str(node["role"]), anchor="ma", fill="#475569", font=small_font)
    elif kind == "system":
        y = max(90, int(height * 0.08))
        for index, node in enumerate(nodes):
            x = int((index + 0.5) * width / node_count)
            draw.rounded_rectangle((x - 130, y, x + 130, y + 78), radius=10, fill="#FFFFFF", outline="#245B73", width=3)
            draw.text((x, y + 22), str(node.get("label", node.get("id", ""))), anchor="ma", fill="#16324F", font=node_font)
            if node.get("role"):
                draw.text((x, y + 49), str(node["role"]), anchor="ma", fill="#475569", font=small_font)
    else:
        y = int(height * 0.52)
        for index, node in enumerate(nodes):
            x = int((index + 0.5) * width / node_count)
            draw.text((x, y), str(node.get("label", node.get("id", ""))), anchor="mm", fill="#16324F", font=node_font)
            if node.get("role"):
                draw.text((x, y + node_font.size + 5), str(node["role"]), anchor="ma", fill="#475569", font=small_font)
    key_text = "关键标识：" + "；".join(str(value) for value in figure.get("labels", []))
    draw.text((35, height - max(35, height // 18)), key_text, fill="#475569", font=small_font)
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, format="PNG", optimize=True)


def _output_manifest(data: Mapping[str, Any], output_dir: Path, figures: Sequence[Mapping[str, Any]]) -> Dict[str, Any]:
    files = []
    for path in sorted(output_dir.rglob("*")):
        if path.is_file() and path.name != "manifest.json":
            files.append({"path": str(path.relative_to(output_dir)), "sha256": _sha256(path), "bytes": path.stat().st_size})
    evidence_summary = {tag: 0 for tag in sorted(EVIDENCE_TYPES)}
    for item in data.get("evidence", []):
        evidence_summary[item["type"]] += 1
    return {
        "manifest_version": "1.0",
        "package_title": data["metadata"]["title"],
        "release_status": data["metadata"]["release_status"],
        "delivery_mode": data["metadata"]["delivery_mode"],
        "route": data["metadata"]["route"],
        "docx_profile": data["metadata"].get("docx_profile", "native"),
        "figure_output_mode": data["metadata"].get("figure_output_mode", "rendered"),
        "reference_docx_sha256": data.get("_reference_docx_sha256"),
        "mode_contract": data.get("mode_contract", {}),
        "requested_sections": data.get("requested_sections", []),
        "rewrite_audit": data.get("rewrite_audit", []),
        "input_sha256": hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest(),
        "sections": [{"id": section["id"], "title": section["title"]} for section in data["sections"]],
        "evidence_summary": evidence_summary,
        "figures": list(figures),
        "audit": data["audit"],
        "files": files,
    }


def render_package(input_path: Path, output_dir: Path, reference_docx: Path | None = None) -> Dict[str, Any]:
    data = load_json(input_path)
    ensure_valid(data)
    profile = data["metadata"].get("docx_profile", "native")
    if profile in {"sample_fidelity", "sample_exact"}:
        resolved_reference = reference_docx or _find_reference_docx(input_path, data["metadata"].get("reference_docx"))
        if resolved_reference is None:
            raise PackageValidationError(["sample_fidelity/sample_exact 找不到 reference_docx"])
        data = dict(data)
        data["_reference_docx_sha256"] = _sha256(resolved_reference)
        reference_docx = resolved_reference
    _check_output_dir(output_dir, data)
    output_dir.mkdir(parents=True, exist_ok=True)
    figure_dir = output_dir / "figures"
    _validate_source_images(data, input_path.parent)
    if data["metadata"].get("figure_output_mode", "rendered") == "spec_only":
        figures = render_figure_specs(data, output_dir)
    else:
        figures = render_figures(data, figure_dir, input_path.parent)
    render_markdown(data, output_dir / "disclosure.md")
    render_docx(data, output_dir / "disclosure.docx", figure_dir, reference_docx)
    manifest = _output_manifest(data, output_dir, figures)
    (output_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return manifest


def _validate_source_images(data: Mapping[str, Any], source_root: Path) -> None:
    if data["metadata"].get("figure_output_mode", "rendered") not in FIGURE_OUTPUT_MODES:
        return
    for figure in data.get("figures", []):
        source_image = figure.get("source_image")
        if not source_image:
            continue
        source_path = (source_root / source_image).resolve()
        source_root_resolved = source_root.resolve()
        if source_root_resolved not in source_path.parents or not source_path.is_file():
            raise PackageValidationError([f"生成图片底图不存在或越界：{source_image}"])


def _check_output_dir(output_dir: Path, data: Mapping[str, Any]) -> None:
    """Reject stale or unexpected files instead of silently mixing packages."""
    if output_dir.is_symlink():
        raise PackageValidationError([f"拒绝向 symlink 输出目录写入：{output_dir}"])
    if not output_dir.exists():
        return
    expected = {"disclosure.md", "disclosure.docx", "manifest.json"}
    if data.get("metadata", {}).get("figure_output_mode", "rendered") == "spec_only":
        expected.add("figure-plan.md")
        expected_figures = set()
    else:
        expected_figures = {
            filename
            for figure in data.get("figures", [])
            for filename in (f"{figure['id']}.svg", f"{figure['id']}.png", f"{figure['id']}-source.png")
            if figure.get("source_image") or not filename.endswith("-source.png")
        }
    for path in output_dir.rglob("*"):
        relative = path.relative_to(output_dir)
        if path.is_symlink():
            raise PackageValidationError([f"输出目录包含 symlink：{relative}"])
        if path.is_dir() and relative == Path("figures"):
            continue
        if path.is_dir():
            raise PackageValidationError([f"输出目录包含未知子目录：{relative}"])
        if relative.parent == Path("figures"):
            if relative.name not in expected_figures:
                raise PackageValidationError([f"输出目录包含残留附图：{relative}"])
        elif relative.parent == Path("."):
            if relative.name not in expected:
                raise PackageValidationError([f"输出目录包含未知文件：{relative}"])
        else:
            raise PackageValidationError([f"输出目录包含越界文件：{relative}"])


def validate_outputs(output_dir: Path) -> List[str]:
    errors: List[str] = []
    manifest_path = output_dir / "manifest.json"
    if not manifest_path.is_file():
        return ["缺少 manifest.json"]
    try:
        manifest = load_json(manifest_path)
    except (OSError, json.JSONDecodeError, PackageValidationError) as exc:
        return [f"manifest 不可读取：{exc}"]
    required_files = ["disclosure.md", "disclosure.docx"]
    if manifest.get("figure_output_mode", "rendered") == "spec_only":
        required_files.append("figure-plan.md")
    for required in required_files:
        if not (output_dir / required).is_file():
            errors.append(f"缺少输出文件：{required}")
    figure_entries = manifest.get("figures", [])
    if not isinstance(figure_entries, list) or any(not isinstance(figure, Mapping) for figure in figure_entries):
        errors.append("manifest.figures 必须是对象数组")
        figure_entries = []
    if manifest.get("delivery_mode") == "FULL_DISCLOSURE" and len(figure_entries) < 3:
        errors.append("manifest.figures 少于 3 张")
    spec_only = manifest.get("figure_output_mode", "rendered") == "spec_only"
    for figure in figure_entries:
        if spec_only:
            if not all(isinstance(figure.get(key), str) and figure.get(key) for key in ("id", "generation_method", "prompt", "source_image")):
                errors.append("manifest figure spec 缺少 id/generation_method/prompt/source_image")
        else:
            if not all(isinstance(figure.get(key), str) and figure.get(key) for key in ("id", "svg", "png")):
                errors.append("manifest figure 缺少 id/svg/png")
                continue
            for key in ("svg", "png"):
                candidate = output_dir / figure[key]
                if not _is_safe_output_path(output_dir, candidate) or not candidate.is_file() or candidate.is_symlink():
                    errors.append(f"附图文件不存在：{figure.get('id')} / {key}")
            if figure.get("source_png"):
                source_candidate = output_dir / figure["source_png"]
                if not _is_safe_output_path(output_dir, source_candidate) or not source_candidate.is_file() or source_candidate.is_symlink():
                    errors.append(f"生成图片底图不存在：{figure.get('id')}")
    try:
        markdown = (output_dir / "disclosure.md").read_text(encoding="utf-8") if (output_dir / "disclosure.md").is_file() else ""
    except (OSError, UnicodeDecodeError) as exc:
        errors.append(f"Markdown 不可读取：{exc}")
        markdown = ""
    for figure in figure_entries:
        if figure.get("id") not in markdown:
            errors.append(f"正文未回指附图：{figure.get('id')}")
    if spec_only:
        try:
            plan_text = (output_dir / "figure-plan.md").read_text(encoding="utf-8")
            for figure in figure_entries:
                if figure.get("id") not in plan_text:
                    errors.append(f"附图方案未回指 figure_id：{figure.get('id')}")
        except (OSError, UnicodeDecodeError) as exc:
            errors.append(f"附图方案不可读取：{exc}")
    manifest_files = manifest.get("files", [])
    if not isinstance(manifest_files, list) or any(not isinstance(entry, Mapping) for entry in manifest_files):
        errors.append("manifest.files 必须是对象数组")
        manifest_files = []
    listed_paths = set()
    for entry in manifest_files:
        path_value = entry.get("path")
        if not isinstance(path_value, str):
            errors.append("manifest.files.path 必须是字符串")
            continue
        candidate = output_dir / path_value
        if not _is_safe_output_path(output_dir, candidate) or candidate.is_symlink() or not candidate.is_file():
            errors.append(f"manifest 文件不存在或路径非法：{path_value}")
            continue
        listed_paths.add(Path(path_value).as_posix())
        expected_hash = entry.get("sha256")
        expected_bytes = entry.get("bytes")
        if not isinstance(expected_hash, str) or _sha256(candidate) != expected_hash:
            errors.append(f"manifest SHA-256 不匹配：{path_value}")
        if not isinstance(expected_bytes, int) or candidate.stat().st_size != expected_bytes:
            errors.append(f"manifest 字节数不匹配：{path_value}")
    actual_paths = {
        path.relative_to(output_dir).as_posix()
        for path in output_dir.rglob("*")
        if path.is_file() and path.name != "manifest.json"
    }
    if listed_paths != actual_paths:
        errors.append(f"manifest 文件清单与实际产物不一致：listed={sorted(listed_paths)}, actual={sorted(actual_paths)}")

    docx_path = output_dir / "disclosure.docx"
    if docx_path.is_file():
        try:
            with zipfile.ZipFile(docx_path) as archive:
                names = set(archive.namelist())
                media = sorted(name for name in names if name.startswith("word/media/"))
                expected_media = 0 if spec_only or manifest.get("docx_profile") == "sample_exact" else len(figure_entries)
                if len(media) != expected_media:
                    errors.append(f"DOCX 媒体数量 {len(media)} 与附图数量 {len(figure_entries)} 不一致")
                document_xml = archive.read("word/document.xml").decode("utf-8")
                if manifest.get("docx_profile") != "sample_exact" and "disclosure" not in document_xml and "技术交底" not in document_xml:
                    errors.append("DOCX 正文缺少交底书文本")
                if manifest.get("docx_profile") not in {"sample_exact", "sample_fidelity"} and "comments.xml" in names:
                    errors.append("DOCX 不应携带审阅批注")
        except (OSError, zipfile.BadZipFile, KeyError) as exc:
            errors.append(f"DOCX 包不可读取：{exc}")
    return errors


def _is_safe_output_path(root: Path, candidate: Path) -> bool:
    try:
        candidate.relative_to(root)
    except ValueError:
        return False
    return not any(part in {"", ".", ".."} for part in candidate.relative_to(root).parts)


def validate_output_or_input(path: Path) -> List[str]:
    if path.is_file():
        try:
            return validate_package(load_json(path))
        except (OSError, json.JSONDecodeError, PackageValidationError, TypeError, ValueError) as exc:
            return [str(exc)]
    if path.is_dir():
        try:
            return validate_outputs(path)
        except (OSError, UnicodeDecodeError, TypeError, ValueError) as exc:
            return [f"产物校验失败：{exc}"]
    return [f"路径不存在：{path}"]
